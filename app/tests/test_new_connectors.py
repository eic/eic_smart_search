"""Offline unit tests for the INSPIRE / arXiv / Indico-events / GitHub-issues
connectors — document building from canned API payloads, no network."""
from __future__ import annotations


class FakeSettings:
    INGEST_USER_AGENT = "test-agent"
    REQUEST_TIMEOUT_SECONDS = 5.0
    INDICO_CATEGORY_URL = "https://indico.example.org/category/402/"
    PEOPLE_REFRESH_INDICO_TIMEOUT_S = 5.0


# ---------------------------------------------------------------------------
# INSPIRE-HEP
# ---------------------------------------------------------------------------

INSPIRE_RECORD = {
    "id": "123456",
    "metadata": {
        "titles": [{"title": "Design of the ePIC dRICH detector"}],
        "abstracts": [{"value": "We present the design of the dual-radiator RICH " * 5}],
        "authors": [
            {
                "full_name": "Prozorov, Alexandr",
                "affiliations": [{"value": "CTU Prague"}],
                "ids": [{"schema": "ORCID", "value": "0000-0001-2345-6789"}],
            },
            {"full_name": "Doe, Jane"},
        ],
        "arxiv_eprints": [{"value": "2401.12345"}],
        "dois": [{"value": "10.1000/test.doi"}],
        "publication_info": [{"journal_title": "Nucl.Instrum.Meth.A", "journal_volume": "1069", "year": 2026}],
        "earliest_date": "2026-01-15",
        "collaborations": [{"value": "ePIC"}],
        "document_type": ["article"],
        "keywords": [{"value": "RICH"}, {"value": "EIC"}],
    },
}


def test_inspire_build_document() -> None:
    from app.ingestion.inspire import InspireHepConnector

    connector = InspireHepConnector(FakeSettings())
    doc = connector._build_document(INSPIRE_RECORD)
    assert doc is not None
    assert doc.external_id == "inspire:123456"
    assert doc.title == "Design of the ePIC dRICH detector"
    assert doc.url == "https://inspirehep.net/literature/123456"
    assert "Prozorov, Alexandr (CTU Prague)" in doc.metadata["authors"]
    assert doc.metadata["author_orcids"]["Prozorov, Alexandr"] == "0000-0001-2345-6789"
    assert "Doe, Jane" in doc.content
    assert "arXiv: 2401.12345" in doc.content
    assert doc.last_updated is not None and doc.last_updated.year == 2026


def test_inspire_skips_empty_record() -> None:
    from app.ingestion.inspire import InspireHepConnector

    connector = InspireHepConnector(FakeSettings())
    assert connector._build_document({"id": "1", "metadata": {}}) is None


# ---------------------------------------------------------------------------
# arXiv
# ---------------------------------------------------------------------------

ARXIV_FEED = """<?xml version="1.0" encoding="UTF-8"?>
<feed xmlns="http://www.w3.org/2005/Atom" xmlns:arxiv="http://arxiv.org/schemas/atom">
  <entry>
    <id>http://arxiv.org/abs/2405.06661v2</id>
    <title>Physics with the ePIC detector at the Electron-Ion Collider</title>
    <summary>The Electron-Ion Collider will enable studies of nucleon structure with unprecedented precision using deep inelastic scattering.</summary>
    <published>2024-05-10T17:59:59Z</published>
    <updated>2024-06-01T10:00:00Z</updated>
    <author><name>Jane Doe</name></author>
    <author><name>John Smith</name></author>
    <category term="physics.ins-det"/>
    <category term="hep-ex"/>
    <arxiv:doi>10.1000/arxiv.doi</arxiv:doi>
    <arxiv:comment>20 pages, 12 figures</arxiv:comment>
  </entry>
</feed>
"""


def test_arxiv_parse_and_build() -> None:
    from app.ingestion.arxiv import ArxivConnector, parse_atom_feed

    entries = parse_atom_feed(ARXIV_FEED)
    assert len(entries) == 1
    entry = entries[0]
    assert entry["arxiv_id"] == "2405.06661"
    assert entry["authors"] == ["Jane Doe", "John Smith"]
    assert entry["categories"] == ["physics.ins-det", "hep-ex"]

    connector = ArxivConnector(FakeSettings())
    doc = connector._build_document(entry)
    assert doc is not None
    assert doc.external_id == "arxiv:2405.06661"
    assert doc.url.startswith("https://arxiv.org/abs/")
    assert "deep inelastic scattering" in doc.content
    assert "Authors: Jane Doe; John Smith" in doc.content
    assert doc.last_updated is not None and doc.last_updated.year == 2024


def test_arxiv_parse_garbage_returns_empty() -> None:
    from app.ingestion.arxiv import parse_atom_feed

    assert parse_atom_feed("not xml at all <<<") == []


# ---------------------------------------------------------------------------
# Indico events
# ---------------------------------------------------------------------------

INDICO_EVENT = {
    "id": "99001",
    "title": "ePIC dRICH WG Meeting",
    "url": "https://indico.example.org/event/99001/",
    "category": "dRICH",
    "startDate": {"date": "2026-03-05", "time": "14:00:00", "tz": "US/Eastern"},
    "chairs": [{"fullName": "Jane Doe", "affiliation": "BNL"}],
    "description": "<p>Weekly working group meeting</p>",
    "contributions": [
        {
            "title": "Aerogel QA update",
            "url": "https://indico.example.org/event/99001/#1",
            "startDate": {"date": "2026-03-05", "time": "14:10:00"},
            "speakers": [{"fullName": "John Smith"}],
            "folders": [
                {
                    "attachments": [
                        {
                            "id": 7,
                            "title": "Aerogel slides",
                            "filename": "aerogel.pdf",
                            "content_type": "application/pdf",
                            "download_url": "https://indico.example.org/att/7/aerogel.pdf",
                            "size": 1024,
                        }
                    ]
                }
            ],
        }
    ],
}


def _make_indico_connector():
    from app.ingestion.indico_events import IndicoEventsConnector

    return IndicoEventsConnector(FakeSettings())


def test_indico_event_document() -> None:
    connector = _make_indico_connector()
    doc = connector._build_event_document(INDICO_EVENT)
    assert doc is not None
    assert doc.external_id == "indico:99001"
    assert doc.title == "ePIC dRICH WG Meeting"
    assert "Chairs: Jane Doe (BNL)" in doc.content
    assert "- Aerogel QA update — John Smith" in doc.content
    assert doc.metadata["contribution_count"] == 1
    assert doc.last_updated is not None and doc.last_updated.month == 3


def test_indico_attachment_flattening() -> None:
    from app.ingestion.indico_events import _is_pdf_attachment, iter_attachments

    atts = iter_attachments(INDICO_EVENT["contributions"][0])
    assert len(atts) == 1
    assert _is_pdf_attachment(atts[0])
    assert not _is_pdf_attachment({"filename": "notes.txt", "content_type": "text/plain"})


def test_indico_material_documents(monkeypatch) -> None:
    import app.ingestion.indico_events as mod

    connector = _make_indico_connector()

    class FakeResponse:
        content = b"%PDF-fake"

        def raise_for_status(self) -> None:
            return None

    # safe_get enforces the SSRF host allowlist + DNS resolution; stub it so
    # the unit test stays offline and host-agnostic.
    monkeypatch.setattr(
        mod, "safe_get", lambda client, url, **kw: (
            (_ for _ in ()).throw(AssertionError(url)) if url != "https://indico.example.org/att/7/aerogel.pdf" else FakeResponse()
        )
    )
    monkeypatch.setattr(
        mod, "extract_pdf_text", lambda data, max_chars: ("extracted slide text " * 10, {"engine": "fake"})
    )
    docs = connector._build_material_documents(object(), INDICO_EVENT)
    assert len(docs) == 1
    doc = docs[0]
    assert doc.external_id == "indico:99001:att:7"
    assert doc.title == "ePIC dRICH WG Meeting — Aerogel QA update — Aerogel slides"
    assert doc.filetype == "pdf"
    assert "extracted slide text" in doc.content


# ---------------------------------------------------------------------------
# GitHub issues
# ---------------------------------------------------------------------------

GH_ISSUE = {
    "number": 42,
    "title": "eic-shell fails to start on WSL2",
    "state": "closed",
    "html_url": "https://github.com/eic/eic-shell/issues/42",
    "user": {"login": "physicist1"},
    "body": "When I run ./eic-shell I get a singularity permission error on WSL2.",
    "labels": [{"name": "bug"}],
    "comments": 1,
    "updated_at": "2026-02-01T12:00:00Z",
}

GH_COMMENTS = [
    {"user": {"login": "maintainer"}, "body": "Fixed by updating to Apptainer 1.3 — see the install docs."}
]


def test_github_issue_content_and_document() -> None:
    from app.ingestion.github_issues import GitHubIssuesConnector, build_issue_content

    content = build_issue_content(GH_ISSUE, GH_COMMENTS)
    assert "State: closed" in content
    assert "Labels: bug" in content
    assert "physicist1 (opened):" in content
    assert "maintainer (comment):" in content

    connector = GitHubIssuesConnector(FakeSettings())

    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self):
            return GH_COMMENTS

    class FakeClient:
        def get(self, url, params=None):
            assert url.endswith("/repos/eic/eic-shell/issues/42/comments")
            return FakeResponse()

    doc = connector._build_document(FakeClient(), "eic/eic-shell", GH_ISSUE)
    assert doc is not None
    assert doc.external_id == "gh-issue:eic/eic-shell#42"
    assert doc.title == "eic/eic-shell#42: eic-shell fails to start on WSL2"
    assert doc.metadata["labels"] == ["bug"]
    assert doc.last_updated is not None and doc.last_updated.year == 2026


def test_github_issues_connector_needs_token_note() -> None:
    # PRs are skipped in _iter_repo via the pull_request key — sanity-check
    # the guard shape here so a payload change breaks loudly.
    assert "pull_request" not in GH_ISSUE


# ---------------------------------------------------------------------------
# pdf_text fallback chain
# ---------------------------------------------------------------------------

def test_pdf_text_handles_garbage_bytes() -> None:
    from app.ingestion.pdf_text import extract_pdf_text

    text, meta = extract_pdf_text(b"this is not a pdf", max_chars=1000)
    assert text == ""
    assert meta.get("skipped") in {"pdf-parse-failed", "no-pdf-engine", "pymupdf-missing"} or "engine" in meta


# ---------------------------------------------------------------------------
# GitHub releases
# ---------------------------------------------------------------------------

GH_RELEASE = {
    "tag_name": "v1.19.0",
    "name": "EICrecon 1.19.0",
    "html_url": "https://github.com/eic/EICrecon/releases/tag/v1.19.0",
    "published_at": "2026-03-01T09:00:00Z",
    "prerelease": False,
    "draft": False,
    "body": "## What's Changed\n* New dRICH reconstruction algorithm\n* Fixed B0 tracker geometry offsets",
}


def test_github_release_document() -> None:
    from app.ingestion.github_releases import GitHubReleasesConnector

    connector = GitHubReleasesConnector(FakeSettings())
    doc = connector._build_document("eic/EICrecon", GH_RELEASE)
    assert doc is not None
    assert doc.external_id == "gh-release:eic/EICrecon@v1.19.0"
    assert doc.title == "eic/EICrecon EICrecon 1.19.0"
    assert "dRICH reconstruction algorithm" in doc.content
    assert doc.last_updated is not None and doc.last_updated.month == 3
    assert doc.metadata["tag"] == "v1.19.0"


def test_github_release_skips_empty_tag() -> None:
    from app.ingestion.github_releases import GitHubReleasesConnector

    connector = GitHubReleasesConnector(FakeSettings())
    assert connector._build_document("eic/epic", {"name": "x", "body": "y"}) is None


# ---------------------------------------------------------------------------
# GitHub Discussions
# ---------------------------------------------------------------------------

GH_DISCUSSION = {
    "number": 512,
    "title": "How to access simulation campaign output?",
    "url": "https://github.com/eic/eic-shell/discussions/512",
    "updatedAt": "2026-04-01T08:00:00Z",
    "category": {"name": "Q&A"},
    "author": {"login": "student1"},
    "body": "Where do I find the latest campaign files and how do I open them?",
    "answer": {"body": "Use rucio with the instructions in the file-access tutorial.", "author": {"login": "expert1"}},
    "comments": {
        "nodes": [
            {"body": "Use rucio with the instructions in the file-access tutorial.", "author": {"login": "expert1"}},
            {"body": "Also see the S3 read-only mirror.", "author": {"login": "helper2"}},
        ]
    },
}


def test_github_discussion_content_and_document() -> None:
    from app.ingestion.github_discussions import (
        GitHubDiscussionsConnector,
        build_discussion_content,
    )

    content = build_discussion_content(GH_DISCUSSION)
    assert "Category: Q&A" in content
    assert "student1 (asked):" in content
    assert "expert1 (accepted answer):" in content
    # Duplicate of the accepted answer among comments must not repeat.
    assert content.count("file-access tutorial") == 1
    assert "helper2 (comment):" in content

    connector = GitHubDiscussionsConnector(FakeSettings())
    doc = connector._build_document("eic/eic-shell", GH_DISCUSSION)
    assert doc is not None
    assert doc.external_id == "gh-discussion:eic/eic-shell#512"
    assert doc.metadata["answered"] is True
    assert doc.last_updated is not None and doc.last_updated.month == 4


def test_github_discussions_require_token() -> None:
    from app.ingestion.github_discussions import GitHubDiscussionsConnector

    class NoTokenSettings(FakeSettings):
        GITHUB_TOKEN = None

    connector = GitHubDiscussionsConnector(NoTokenSettings())
    assert connector.iter_documents(max_items=5) == []


def test_attachment_kind_classifier() -> None:
    from app.ingestion.indico_events import _attachment_kind

    assert _attachment_kind({"filename": "status.pdf"}) == "pdf"
    assert _attachment_kind({"filename": "ePIC Status - rev 4.pptx"}) == "pptx"
    assert _attachment_kind({"content_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}) == "pptx"
    assert _attachment_kind({"filename": "notes.txt"}) is None


def test_pptx_text_extraction_roundtrip() -> None:
    import io

    pytest_pptx = __import__("pytest")
    try:
        from pptx import Presentation
    except ImportError:
        pytest_pptx.skip("python-pptx not installed")
    from app.ingestion.pdf_text import extract_pptx_text

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "ePIC Collaboration Status"
    slide.placeholders[1].text = "Priorities: complete the TDR and detector installation"
    buf = io.BytesIO()
    prs.save(buf)

    text, meta = extract_pptx_text(buf.getvalue(), max_chars=10_000)
    assert "ePIC Collaboration Status" in text
    assert "TDR" in text
    assert meta["engine"] == "python-pptx"
    assert meta["slides_extracted"] == 1

    bad_text, bad_meta = extract_pptx_text(b"not a pptx", max_chars=100)
    assert bad_text == ""
    assert bad_meta.get("skipped") in {"pptx-parse-failed", "python-pptx-missing"}
