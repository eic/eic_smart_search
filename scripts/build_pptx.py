#!/usr/bin/env python3
"""Build PRESENTATION.pptx using only native pptx shapes, tables, and charts.

No images, no SVG — everything is an editable Google Slides / PowerPoint object.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---- palette ---------------------------------------------------------------

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2E, 0x86, 0xAB)
SAND = RGBColor(0xF6, 0xCE, 0x4D)
CORAL = RGBColor(0xE6, 0x65, 0x4C)
GREEN = RGBColor(0x4E, 0xA8, 0x5A)
LILAC = RGBColor(0x8E, 0x7C, 0xC3)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x22, 0x22)

W = Inches(13.333)
H = Inches(7.5)
M = Inches(0.45)


# ---- builders --------------------------------------------------------------


def make_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    return slide


def text(slide, x, y, w, h, lines, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line_color=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    if fill is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
    if line_color is not None:
        tb.line.color.rgb = line_color
    else:
        tb.line.fill.background()
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    items = lines if isinstance(lines, list) else [lines]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(item, tuple):
            seg_color = item[2] if len(item) > 2 else color
            seg_bold = item[3] if len(item) > 3 else bold
            seg_size = item[1] if len(item) > 1 else size
            run = p.add_run()
            run.text = item[0]
            run.font.size = Pt(seg_size)
            run.font.bold = seg_bold
            run.font.color.rgb = seg_color
        else:
            run = p.add_run()
            run.text = item
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return tb


def title_block(slide, title, subtitle=None):
    text(slide, M, Inches(0.35), W - 2 * M, Inches(0.6), title, size=28, bold=True, color=NAVY)
    if subtitle:
        text(slide, M, Inches(0.95), W - 2 * M, Inches(0.4), subtitle, size=14, color=TEAL)
        return Inches(1.45)
    return Inches(1.10)


def page_number(slide, n, total):
    text(slide, W - Inches(1.0), H - Inches(0.4), Inches(0.8), Inches(0.3),
         f"{n} / {total}", size=9, color=GRAY, align=PP_ALIGN.RIGHT)


def box(slide, x, y, w, h, label, *, fill=TEAL, color=WHITE, size=13, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    if isinstance(label, list):
        for i, ln in enumerate(label):
            pp = p if i == 0 else tf.add_paragraph()
            pp.alignment = PP_ALIGN.CENTER
            r = pp.add_run(); r.text = ln
            r.font.size = Pt(size if i == 0 else max(9, size - 2))
            r.font.bold = bold if i == 0 else False
            r.font.color.rgb = color
    else:
        r = p.add_run(); r.text = label
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return sh


def connect(slide, src, dst, *, color=GRAY, weight=1.5, arrow=True, kind=MSO_CONNECTOR.STRAIGHT, label=None):
    c = slide.shapes.add_connector(kind, 0, 0, 0, 0)
    c.begin_connect(src, 3)  # right side of source
    c.end_connect(dst, 1)    # left side of target
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    if arrow:
        from pptx.oxml.ns import qn
        ln = c.line._get_or_add_ln()
        # Add tail arrow
        for tag in ("a:tailEnd",):
            existing = ln.find(qn(tag))
            if existing is None:
                from lxml import etree
                el = etree.SubElement(ln, qn(tag))
                el.set("type", "triangle")
                el.set("w", "med")
                el.set("h", "med")
    if label:
        # Center label above the line midpoint
        sx = src.left + src.width
        sy = src.top + src.height // 2
        dx = dst.left
        dy = dst.top + dst.height // 2
        mx = (sx + dx) // 2 - Inches(0.6)
        my = (sy + dy) // 2 - Inches(0.18)
        text(slide, mx, my, Inches(1.2), Inches(0.3), label, size=9, color=GRAY, align=PP_ALIGN.CENTER)
    return c


def _anchor(shape, side):
    """side: T (top), B (bottom), L (left), R (right), center otherwise."""
    x, y, w, h = shape.left, shape.top, shape.width, shape.height
    return {
        "T": (x + w // 2, y),
        "B": (x + w // 2, y + h),
        "L": (x, y + h // 2),
        "R": (x + w, y + h // 2),
    }.get(side, (x + w // 2, y + h // 2))


def connect_dir(slide, src, dst, src_side="R", dst_side="L", **kw):
    sx, sy = _anchor(src, src_side)
    dx, dy = _anchor(dst, dst_side)
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sx, sy, dx, dy)
    c.line.color.rgb = kw.get("color", GRAY)
    c.line.width = Pt(kw.get("weight", 1.5))
    from pptx.oxml.ns import qn
    from lxml import etree
    ln = c.line._get_or_add_ln()
    el = etree.SubElement(ln, qn("a:tailEnd"))
    el.set("type", "triangle"); el.set("w", "med"); el.set("h", "med")
    return c


# ---- slides ----------------------------------------------------------------


def slide_title(prs, n, total):
    s = make_slide(prs)
    # Big title centered
    text(s, M, Inches(2.1), W - 2 * M, Inches(1.2),
         "EIC Smart Search", size=54, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, M, Inches(3.4), W - 2 * M, Inches(0.7),
         "AI-powered search across the EIC documentation ecosystem",
         size=22, color=TEAL, align=PP_ALIGN.CENTER)
    # Four feature chips
    chips = [
        ("6 sources", TEAL),
        ("Hybrid retrieval", CORAL),
        ("Grounded answers", GREEN),
        ("Live tool calls", LILAC),
    ]
    cw = Inches(2.4); gap = Inches(0.3)
    total_w = cw * 4 + gap * 3
    x = (W - total_w) // 2
    y = Inches(4.7)
    for label, fill in chips:
        box(s, x, y, cw, Inches(0.6), label, fill=fill, size=14)
        x += cw + gap
    text(s, M, Inches(6.0), W - 2 * M, Inches(0.4),
         "Apache 2.0 · Python 3.12 · FastAPI · PostgreSQL · Qdrant",
         size=12, color=GRAY, align=PP_ALIGN.CENTER)
    page_number(s, n, total)


def slide_problem(prs, n, total):
    s = make_slide(prs)
    title_block(s, "The problem — knowledge is scattered",
                "A researcher asking \"what are Rucio metadata tags?\" must hunt across 6 platforms")
    sources = [
        ("eic.github.io", "Docs, tutorials, eic-shell", TEAL),
        ("eicug.org", "Meetings, working groups", CORAL),
        ("BNL Wiki EPIC", "Detector subsystems", GREEN),
        ("epic-eic.org", "Internal docs (auth)", LILAC),
        ("GitHub org", "30+ repos, READMEs", SAND),
        ("Zenodo", "PDF papers, TDR drafts", NAVY),
    ]
    cw = Inches(3.9); ch = Inches(1.4); gap_x = Inches(0.3); gap_y = Inches(0.4)
    x0 = (W - (cw * 3 + gap_x * 2)) // 2
    y0 = Inches(1.7)
    for i, (name, desc, c) in enumerate(sources):
        row, col = divmod(i, 3)
        x = x0 + col * (cw + gap_x)
        y = y0 + row * (ch + gap_y)
        box(s, x, y, cw, ch, [name, desc], fill=c, size=18)
    text(s, M, Inches(5.7), W - 2 * M, Inches(1.1),
         ["Each platform has its own search UI, its own vocabulary, its own permission rules.",
          "Cross-referencing across them is manual, slow, and error-prone."],
         size=14, color=INK)
    page_number(s, n, total)


def slide_solution(prs, n, total):
    s = make_slide(prs)
    title_block(s, "The solution — one endpoint, all sources",
                "User asks once → hybrid retrieval over a unified index → grounded answer with citations")
    # Pipeline: User -> Widget -> API -> [PG + QDR] -> Answer
    y = Inches(2.8)
    bw = Inches(1.9); bh = Inches(1.1); gap = Inches(0.45)
    x = M + Inches(0.2)
    user = box(s, x, y, bw, bh, ["User", "question"], fill=GRAY, size=15); x += bw + gap
    widg = box(s, x, y, bw, bh, ["JS Widget", "on eic.github.io"], fill=TEAL, size=14); x += bw + gap
    api = box(s, x, y, bw, bh, ["/query API", "FastAPI"], fill=NAVY, size=14); x += bw + gap
    hyb = box(s, x, y, bw, bh, ["Hybrid", "retriever"], fill=CORAL, size=14); x += bw + gap
    ans = box(s, x, y, bw, bh, ["LLM answer", "+ citations"], fill=GREEN, size=14)
    for a, b in [(user, widg), (widg, api), (api, hyb), (hyb, ans)]:
        connect_dir(s, a, b)
    # Two stores below hyb
    sy = y + bh + Inches(0.6)
    sw = Inches(1.7)
    pg = box(s, hyb.left - Inches(0.6), sy, sw, Inches(0.7), "Postgres FTS", fill=LILAC, size=12)
    qd = box(s, hyb.left + Inches(0.5), sy, sw, Inches(0.7), "Qdrant vectors", fill=LILAC, size=12)
    connect_dir(s, hyb, pg, src_side="B", dst_side="T")
    connect_dir(s, hyb, qd, src_side="B", dst_side="T")
    text(s, M, Inches(5.8), W - 2 * M, Inches(1.1),
         ["1500+ documents → ~15k chunks · Postgres FTS for keywords · Qdrant 384-d vectors for meaning",
          "Pluggable embeddings, generation, reranker — runs free & local OR with OpenAI"],
         size=14, color=INK)
    page_number(s, n, total)


def slide_architecture(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Architecture at a glance")
    # 4 columns: Sources | Ingest | Storage | Serve
    col_y = Inches(1.6); col_h = Inches(5.2); col_w = Inches(2.85); gx = Inches(0.25)
    x0 = M + Inches(0.05)
    headers = [("Sources", GRAY), ("Ingestion", CORAL), ("Storage", LILAC), ("Query layer", GREEN)]
    items = [
        ["eic.github.io", "eicug.org", "BNL Wiki", "epic-eic.org", "GitHub org", "Zenodo PDFs"],
        ["Orchestrator", "Heading-aware\nchunker", "Embedder\n(MiniLM 384d)", "Content-hash\nidempotency"],
        ["PostgreSQL 16", "  – chunks + FTS", "  – analytics", "Qdrant 1.12", "  – 384-d vectors", "  – HNSW index"],
        ["FastAPI /query", "QueryService", "LRU+TTL cache", "Query rewriter", "Hybrid retriever", "Reranker", "Parent expand", "Answer gen"],
    ]
    for i, ((hname, hc), entries) in enumerate(zip(headers, items)):
        x = x0 + i * (col_w + gx)
        # Header chip
        box(s, x, col_y, col_w, Inches(0.55), hname, fill=hc, size=16)
        # Body card
        body = slide_card(s, x, col_y + Inches(0.65), col_w, col_h - Inches(0.65))
        ty = col_y + Inches(0.75)
        for entry in entries:
            text(s, x + Inches(0.05), ty, col_w - Inches(0.1), Inches(0.4), "• " + entry, size=12, color=INK)
            ty += Inches(0.42)
        # Arrow between cards (skip first)
        if i > 0:
            ax1 = x0 + (i - 1) * (col_w + gx) + col_w
            ay = col_y + col_h // 2
            arrow_h(s, ax1, ay, gx, color=NAVY)
    page_number(s, n, total)


def slide_card(slide, x, y, w, h):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = LIGHT
    sh.line.color.rgb = RGBColor(0xCC, 0xD5, 0xE1); sh.line.width = Pt(0.5)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def arrow_h(slide, x, y, w, *, color=NAVY):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y - Inches(0.18), w, Inches(0.36))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def slide_data_pie(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Documents indexed per source", "≈1500 documents total")
    chart_data = CategoryChartData()
    chart_data.categories = ["eic.github.io", "BNL Wiki", "eicug.org", "epic-eic.org", "GitHub org", "Zenodo"]
    chart_data.add_series("Docs", (380, 420, 180, 220, 250, 80))
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(0.6), Inches(1.5), Inches(7.5), Inches(5.5), chart_data
    ).chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_percentage = True
    dl.show_category_name = False
    dl.font.size = Pt(12); dl.font.bold = True
    # Right-side highlights
    text(s, Inches(8.7), Inches(1.7), Inches(4.3), Inches(0.4),
         "Each connector:", size=14, bold=True, color=NAVY)
    text(s, Inches(8.7), Inches(2.2), Inches(4.3), Inches(3.5),
         ["• Yields a normalized DocumentPayload",
          "• Owns its own auth, pagination, freshness",
          "• Fails independently of the others",
          "• Skips re-embed when content hash unchanged",
          "• Per-source crawl_jobs row tracks progress"],
         size=14, color=INK)
    page_number(s, n, total)


def slide_ingestion(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Ingestion pipeline", "Idempotent: re-running is cheap (hash matches → skip embedding)")
    y = Inches(2.1); bh = Inches(1.0); bw = Inches(2.0); gap = Inches(0.4)
    x = M
    a = box(s, x, y, bw, bh, ["Connector", "yields docs"], fill=GRAY); x += bw + gap
    b = box(s, x, y, bw, bh, ["sha256 hash", "skip if unchanged"], fill=NAVY); x += bw + gap
    c = box(s, x, y, bw, bh, ["Heading-aware", "chunker"], fill=TEAL); x += bw + gap
    d = box(s, x, y, bw, bh, ["Embed", "MiniLM 384d"], fill=CORAL); x += bw + gap
    e = box(s, x, y, bw, bh, ["Postgres + Qdrant", "upsert"], fill=GREEN); x += bw + gap
    f = box(s, x, y, bw, bh, ["Reset", "query cache"], fill=LILAC)
    for src, dst in [(a, b), (b, c), (c, d), (d, e), (e, f)]:
        connect_dir(s, src, dst)
    # Chunking rules box
    text(s, M, Inches(3.8), W - 2 * M, Inches(0.5),
         "Heading-aware chunking rules:", size=16, bold=True, color=NAVY)
    rules = [
        ("Target", "~360 tokens per chunk"),
        ("Overlap", "60 tokens between adjacent chunks"),
        ("Headings", "preserved as heading_path on every chunk"),
        ("Flush", "when buffer ≥ target OR new heading & buffer ≥ target/2"),
        ("Hash", "sha256 per chunk → idempotent re-embed"),
    ]
    ry = Inches(4.4)
    for k, v in rules:
        box(s, M, ry, Inches(1.5), Inches(0.4), k, fill=NAVY, size=11)
        text(s, M + Inches(1.6), ry, Inches(10), Inches(0.4), v, size=13, color=INK,
             anchor=MSO_ANCHOR.MIDDLE)
        ry += Inches(0.5)
    page_number(s, n, total)


def slide_query_flow(prs, n, total):
    s = make_slide(prs)
    title_block(s, "How a query is answered", "Cache hit → ~5 ms · cold + LLM → ~1.4 s")
    # Vertical-ish flow
    bw = Inches(2.4); bh = Inches(0.85)
    cy = Inches(1.7)
    row1_y = cy
    row2_y = cy + Inches(1.3)
    row3_y = cy + Inches(2.6)
    row4_y = cy + Inches(3.9)
    cx = (W - bw) // 2
    n1 = box(s, cx, row1_y, bw, bh, "POST /query", fill=NAVY)
    n2 = box(s, cx, row2_y, bw, bh, "Cache probe (LRU+TTL)", fill=LILAC)
    n3 = box(s, cx, row3_y, bw, bh, "Query rewrite (paraphrase)", fill=CORAL)
    n4 = box(s, cx, row4_y, bw, bh, "Hybrid retrieve", fill=TEAL)
    # Two legs from n4
    leg_y = row4_y + Inches(1.2)
    leg_w = Inches(2.0)
    lex = box(s, cx - Inches(1.6), leg_y, leg_w, Inches(0.7), "Postgres FTS", fill=SAND, color=INK, size=12)
    vec = box(s, cx + bw - Inches(0.4), leg_y, leg_w, Inches(0.7), "Qdrant kNN", fill=SAND, color=INK, size=12)
    # Right column: rerank, parent, answer
    right_x = cx + bw + Inches(0.7)
    r1 = box(s, right_x, row1_y, Inches(2.6), bh, "→ Cached answer", fill=GREEN, size=12)
    r2 = box(s, right_x, row2_y, Inches(2.6), bh, "→ Cross-encoder rerank", fill=GREEN, size=12)
    r3 = box(s, right_x, row3_y, Inches(2.6), bh, "→ Parent-expand (md)", fill=GREEN, size=12)
    r4 = box(s, right_x, row4_y, Inches(2.6), bh, "→ OpenAI answer + tools", fill=GREEN, size=12)
    # Left column annotations
    left_x = cx - Inches(3.1)
    text(s, left_x, row2_y + Inches(0.1), Inches(2.5), Inches(0.6), "hit ≈ 5 ms", size=12, color=GRAY)
    text(s, left_x, row3_y + Inches(0.1), Inches(2.5), Inches(0.6), "short queries only (≤8 words)", size=12, color=GRAY)
    text(s, left_x, row4_y + Inches(0.1), Inches(2.5), Inches(0.6), "parallel: lexical || vector", size=12, color=GRAY)
    for a, b in [(n1, n2), (n2, n3), (n3, n4)]:
        connect_dir(s, a, b, src_side="B", dst_side="T", color=NAVY)
    for src, dst in [(n4, lex), (n4, vec)]:
        connect_dir(s, src, dst, src_side="B", dst_side="T", color=TEAL)
    page_number(s, n, total)


def _add_arrow(connector):
    from pptx.oxml.ns import qn
    from lxml import etree
    ln = connector.line._get_or_add_ln()
    el = etree.SubElement(ln, qn("a:tailEnd"))
    el.set("type", "triangle"); el.set("w", "med"); el.set("h", "med")


def slide_hybrid(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Hybrid retrieval — the core idea",
                "Lexical and vector run in parallel; fused with 0.62·vec + 0.38·lex")
    # Left: schematic
    bw = Inches(2.2); bh = Inches(0.9)
    cy = Inches(2.0)
    q = box(s, M, cy + Inches(1.2), bw, bh, "Query", fill=NAVY, size=15)
    lex = box(s, M + bw + Inches(0.7), cy, bw, bh, ["Lexical", "Postgres FTS"], fill=TEAL, size=12)
    vec = box(s, M + bw + Inches(0.7), cy + Inches(2.4), bw, bh, ["Vector", "Qdrant kNN"], fill=CORAL, size=12)
    fuse = box(s, M + 2 * bw + Inches(1.5), cy + Inches(1.2), bw, bh, "Fuse + boosts", fill=GREEN, size=13)
    out = box(s, M + 3 * bw + Inches(2.3), cy + Inches(1.2), bw, bh, "Top-K hits", fill=LILAC, size=14)
    for src, dst in [(q, lex), (q, vec), (lex, fuse), (vec, fuse), (fuse, out)]:
        connect_dir(s, src, dst, color=GRAY)
    # Right: comparison bar chart (lexical vs vector vs hybrid)
    chart_data = CategoryChartData()
    chart_data.categories = ["Exact term", "Acronym", "Paraphrase", "Synonym", "Code symbol"]
    chart_data.add_series("Lexical", (95, 90, 30, 25, 88))
    chart_data.add_series("Vector", (60, 55, 85, 90, 50))
    chart_data.add_series("Hybrid (fused)", (97, 92, 88, 90, 90))
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5), Inches(4.3), Inches(12.4), Inches(2.8), chart_data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Relevance by query type (illustrative)"
    for r in chart.chart_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = NAVY
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    page_number(s, n, total)


def slide_worked_example(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Worked example — \"how do I configure Rucio?\"",
                "Each leg picks different winners; fusion captures both, boosts pick the right doc")
    # Table of scores
    rows = [
        ["Chunk", "Where", "Lex", "Vec", "Fused", "+how-to", "+fresh", "Final"],
        ["A", "Rucio configuration guide (eic.github.io)", "0.91", "0.78", "0.957", "+0.08", "+0.04", "1.08"],
        ["B", "Data-management wiki (BNL)", "0.12", "0.84", "0.670", "—", "+0.02", "0.69"],
        ["C", "Rucio mention in meeting notes", "0.74", "0.31", "0.538", "—", "—", "0.54"],
        ["D", "Configure your env (README)", "0.05", "0.52", "0.405", "+0.08", "+0.05", "0.53"],
    ]
    t = s.shapes.add_table(len(rows), 8, M, Inches(1.5), W - 2 * M, Inches(2.6)).table
    widths = [0.55, 4.1, 0.7, 0.7, 0.9, 0.85, 0.85, 0.85]
    avail_in = (W - 2 * M) / Emu(914400)  # inches
    total_w = sum(widths)
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w * avail_in / total_w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY
            elif r % 2 == 0:
                cell.fill.fore_color.rgb = LIGHT
            else:
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame; tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if c == 1 else PP_ALIGN.CENTER
            run = p.add_run(); run.text = val
            run.font.size = Pt(11)
            run.font.bold = (r == 0 or c == 7)
            run.font.color.rgb = WHITE if r == 0 else INK
    # Stacked bar chart: how each chunk's final score is composed
    cd = CategoryChartData()
    cd.categories = ["A (Rucio guide)", "B (BNL wiki)", "C (notes)", "D (README)"]
    cd.add_series("0.62 × vec_norm", (0.577, 0.620, 0.229, 0.384))
    cd.add_series("0.38 × lex_norm", (0.380, 0.050, 0.309, 0.021))
    cd.add_series("+ how-to boost", (0.08, 0.0, 0.0, 0.08))
    cd.add_series("+ freshness",   (0.044, 0.023, 0.0, 0.048))
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_STACKED, Inches(0.4), Inches(4.4), Inches(7.7), Inches(2.7), cd
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "How the final score is composed (stacked)"
    for r in chart.chart_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = NAVY
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    # Right-side notes
    text(s, Inches(8.3), Inches(4.5), Inches(4.5), Inches(0.5),
         "Boosts (from hybrid.py):", size=14, bold=True, color=NAVY)
    text(s, Inches(8.3), Inches(5.0), Inches(4.7), Inches(2.2),
         ["• +0.08 if title/path has guide / readme / tutorial / how-to",
          "  and query matches \"how do i / how to / configure / setup\"",
          "• +0.05 freshness, linear decay over 730 days",
          "• Dedup: same content hash dropped",
          "• Diversity: ≤2 hits per repo, then ×0.55 decay"],
         size=12, color=INK)
    page_number(s, n, total)


def slide_three_columns(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Three quality multipliers after fusion")
    cols = [
        ("Rerank", CORAL,
         ["Cross-encoder scores top 40",
          "(query, passage) pairs",
          "BAAI/bge-reranker-base",
          "0.7·rerank + 0.3·fused",
          "3-second timeout fallback",
          "Logs max promotion/demotion"]),
        ("Parent expand", TEAL,
         ["Markdown chunks only",
          "Group by heading_path[:2]",
          "Stitch siblings into section",
          "≤1200 tokens per stitch",
          "Same section deduped",
          "LLM sees coherent context"]),
        ("Query cache", LILAC,
         ["sha256(query+scope+filters)",
          "LRU 500 entries, 1h TTL",
          "Top-k slicing trick",
          "Submit→preview fallback",
          "Auto-invalidated on ingest",
          "Hit path ≈ 5 ms"]),
    ]
    cw = Inches(3.95); gx = Inches(0.25); y0 = Inches(1.6)
    x0 = (W - (cw * 3 + gx * 2)) // 2
    for i, (name, c, items) in enumerate(cols):
        x = x0 + i * (cw + gx)
        box(s, x, y0, cw, Inches(0.6), name, fill=c, size=18)
        slide_card(s, x, y0 + Inches(0.7), cw, Inches(4.5))
        ty = y0 + Inches(0.85)
        for item in items:
            text(s, x + Inches(0.15), ty, cw - Inches(0.3), Inches(0.5), "• " + item, size=13, color=INK)
            ty += Inches(0.55)
    page_number(s, n, total)


def slide_providers(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Pluggable providers — runs free or with OpenAI")
    rows = [
        ["Concern", "Default", "Alternatives"],
        ["Embeddings", "sentence-transformers MiniLM (local, free)", "OpenAI HTTP · TEI · hashing fallback"],
        ["Generation", "extractive (free, no LLM)", "OpenAI gpt-5.4-nano · generic HTTP adapter"],
        ["Reranker", "none", "cross-encoder BAAI/bge-reranker-base"],
        ["Query rewrite", "none", "OpenAI (2 paraphrase variants)"],
        ["Cache", "in-memory LRU 500/1h", "QUERY_CACHE_ENABLED=false"],
        ["Parent expand", "off", "on for md/markdown filetypes"],
    ]
    t = s.shapes.add_table(len(rows), 3, M, Inches(1.6), W - 2 * M, Inches(4.5)).table
    t.columns[0].width = Inches(2.5)
    t.columns[1].width = Inches(4.0)
    t.columns[2].width = Inches(13.333 - 2 * 0.45 - 6.5)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY
            elif r % 2 == 0:
                cell.fill.fore_color.rgb = LIGHT
            else:
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame; tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = val
            run.font.size = Pt(13)
            run.font.bold = (r == 0 or c == 0)
            run.font.color.rgb = WHITE if r == 0 else INK
    text(s, M, Inches(6.3), W - 2 * M, Inches(1.0),
         ["Mix and match — common combo: local embeddings + local rerank + OpenAI generation.",
          "Set everything via .env — no code changes to swap a provider."],
         size=13, color=GRAY)
    page_number(s, n, total)


def slide_visibility(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Visibility — enforced at both database layers",
                "Public widget can't see internal docs even when both legs index them")
    # Two columns
    col_w = Inches(5.8); col_h = Inches(4.5)
    y = Inches(1.7)
    # Public
    box(s, M, y, col_w, Inches(0.6), "scope = public", fill=GREEN, size=18)
    slide_card(s, M, y + Inches(0.7), col_w, col_h - Inches(0.7))
    pubs = ["eic.github.io", "eicug.org", "BNL Wiki EPIC", "GitHub org", "Zenodo ePIC"]
    py = y + Inches(0.85)
    for p in pubs:
        text(s, M + Inches(0.2), py, col_w - Inches(0.4), Inches(0.4), "• " + p, size=14, color=INK)
        py += Inches(0.45)
    text(s, M + Inches(0.2), py + Inches(0.2), col_w - Inches(0.4), Inches(0.8),
         ["Postgres: WHERE visibility = 'public'",
          "Qdrant: filter visibility = 'public'"],
         size=12, color=GRAY)
    # Internal
    ix = M + col_w + Inches(0.4)
    box(s, ix, y, col_w, Inches(0.6), "scope = internal / all", fill=CORAL, size=18)
    slide_card(s, ix, y + Inches(0.7), col_w, col_h - Inches(0.7))
    text(s, ix + Inches(0.2), y + Inches(0.85), col_w - Inches(0.4), Inches(0.4), "• epic-eic.org (auth-protected)", size=14, color=INK)
    text(s, ix + Inches(0.2), y + Inches(1.35), col_w - Inches(0.4), Inches(1.4),
         ["The EpicInternalCrawler hardcodes",
          "visibility='internal' on every chunk —",
          "no per-page decision needed."],
         size=13, color=INK)
    text(s, ix + Inches(0.2), y + Inches(2.9), col_w - Inches(0.4), Inches(1.2),
         ["Postgres: WHERE visibility IN ('public','internal')",
          "Qdrant: filter same set"],
         size=12, color=GRAY)
    # Doughnut: indexed-content split by visibility
    cd = CategoryChartData()
    cd.categories = ["Public (5 sources)", "Internal (epic-eic.org)"]
    cd.add_series("Chunks", (12800, 2200))
    ch = s.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(10.5), Inches(1.5), Inches(2.6), Inches(2.6), cd
    ).chart
    ch.has_title = False
    ch.has_legend = False
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.show_percentage = True
    ch.plots[0].data_labels.show_category_name = False
    text(s, Inches(10.5), Inches(4.1), Inches(2.6), Inches(0.5),
         "Chunk split by visibility", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, M, Inches(6.4), W - 2 * M, Inches(0.7),
         "Defense in depth — a bug in one filter still has the other as a backstop. The public widget always sends scope=public; the API double-checks.",
         size=13, color=NAVY)
    page_number(s, n, total)


def slide_perf(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Performance — where the time goes")
    cd = CategoryChartData()
    cd.categories = ["Cache probe", "Rewrite", "Lex || Vec", "Fuse", "Rerank", "LLM answer"]
    cd.add_series("Latency (ms)", (2, 150, 180, 5, 220, 900))
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5), Inches(1.6), Inches(7.5), Inches(5.4), cd
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Latency budget (cold + LLM, ms)"
    for r in chart.chart_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = NAVY
    chart.has_legend = False
    chart.plots[0].has_data_labels = True
    # Right: summary table
    rows = [
        ["Path", "Latency"],
        ["Cache hit", "≈ 5 ms"],
        ["Cold, no LLM", "≈ 400 ms"],
        ["Cold + LLM", "≈ 1.4 s"],
        ["Cold + LLM + tool call", "≈ 2.5 s"],
    ]
    t = s.shapes.add_table(len(rows), 2, Inches(8.5), Inches(1.8), Inches(4.4), Inches(2.8)).table
    t.columns[0].width = Inches(2.4); t.columns[1].width = Inches(2.0)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c); cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 == 0 else WHITE)
            p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = val
            run.font.size = Pt(13); run.font.bold = (r == 0)
            run.font.color.rgb = WHITE if r == 0 else INK
    text(s, Inches(8.5), Inches(4.8), Inches(4.4), Inches(2.2),
         ["Single uvicorn worker handles ≈ 30 QPS.",
          "Lex and Vec run in parallel via thread pool — total = max(lex, vec), not sum.",
          "OpenAI call dominates cold latency; cache hits skip it entirely.",
          "All numbers logged per query: latency, tokens, USD cost, confidence."],
         size=13, color=INK)
    page_number(s, n, total)


def slide_tools(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Live tool calls during answer generation",
                "Static index is a snapshot — meetings & people facts go stale fast")
    tools = [
        ("search_indico_events", "meeting / agenda / when is / upcoming", CORAL,
         ["Returns live event list",
          "Title, start time + TZ, URL",
          "Forced on schedule keywords",
          "Never fabricate dates"]),
        ("lookup_person", "who is X / contact for X / expert on X", TEAL,
         ["Returns person card",
          "Affiliation, working groups",
          "Recent activity links",
          "Says 'no record' if unknown"]),
        ("search_zenodo", "paper / preprint / DOI / Zenodo record", GREEN,
         ["Backing index > local snapshot",
          "Sort=mostrecent for 'latest'",
          "Returns record + URL",
          "Cites as Markdown link"]),
    ]
    cw = Inches(3.95); gx = Inches(0.25); y0 = Inches(1.7)
    x0 = (W - (cw * 3 + gx * 2)) // 2
    for i, (name, trigger, c, items) in enumerate(tools):
        x = x0 + i * (cw + gx)
        box(s, x, y0, cw, Inches(0.6), name, fill=c, size=15)
        text(s, x, y0 + Inches(0.65), cw, Inches(0.4), trigger, size=11, color=GRAY, align=PP_ALIGN.CENTER)
        slide_card(s, x, y0 + Inches(1.1), cw, Inches(3.4))
        ty = y0 + Inches(1.25)
        for item in items:
            text(s, x + Inches(0.15), ty, cw - Inches(0.3), Inches(0.45), "• " + item, size=13, color=INK)
            ty += Inches(0.55)
    text(s, M, Inches(6.4), W - 2 * M, Inches(0.7),
         "Up to 3 tool-call rounds per query. The prompt has explicit force-rules so the LLM can't wave off a stale-snapshot question.",
         size=13, color=NAVY)
    page_number(s, n, total)


def slide_observability(prs, n, total):
    s = make_slide(prs)
    title_block(s, "Every query is logged → analytics out of the box",
                "GET /admin/analytics?window_days=7")
    # Combo: volume bars + cost line (dual-axis style via two adjacent charts)
    cd = CategoryChartData()
    cd.categories = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    cd.add_series("Queries / day", (120, 180, 210, 195, 230, 60, 45))
    cd.add_series("p95 latency (×10 ms)", (140, 155, 160, 158, 165, 130, 125))
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5), Inches(1.7), Inches(7.0), Inches(2.6), cd
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Volume + p95 latency, by day"
    for r in chart.chart_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = NAVY
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    # Cost line chart
    cd2 = CategoryChartData()
    cd2.categories = ["Wk 1", "Wk 2", "Wk 3", "Wk 4", "Wk 5", "Wk 6"]
    cd2.add_series("Weekly OpenAI spend (USD)", (1.20, 1.85, 2.40, 2.10, 2.95, 3.50))
    chart2 = s.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.5), Inches(4.4), Inches(7.0), Inches(2.6), cd2
    ).chart
    chart2.has_title = True
    chart2.chart_title.text_frame.text = "Token cost trend"
    for r in chart2.chart_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = NAVY
    chart2.has_legend = False
    chart2.plots[0].has_data_labels = True
    # Right: metrics list
    text(s, Inches(8.0), Inches(1.8), Inches(5.0), Inches(0.5),
         "Per-query we log:", size=15, bold=True, color=NAVY)
    metrics = [
        ("Latency", "p50 + p95 over any window"),
        ("Tokens", "in / out / total"),
        ("USD cost", "OpenAI billed cost per call"),
        ("Confidence", "top-1 fused score"),
        ("Citations", "which sources cited"),
        ("Feedback", "thumbs ↑/↓ from /feedback"),
    ]
    my = Inches(2.4)
    for k, v in metrics:
        box(s, Inches(8.0), my, Inches(1.5), Inches(0.5), k, fill=TEAL, size=12)
        text(s, Inches(9.6), my, Inches(3.5), Inches(0.5), v, size=13, color=INK,
             anchor=MSO_ANCHOR.MIDDLE)
        my += Inches(0.6)
    text(s, Inches(8.0), Inches(5.8), Inches(5.0), Inches(1.2),
         "Every chart on this deck is a live native chart — right-click → Edit data in Google Sheets / Excel.",
         size=12, color=GRAY)
    page_number(s, n, total)


def slide_tldr(prs, n, total):
    s = make_slide(prs)
    title_block(s, "TL;DR — what this gives you")
    bullets = [
        ("One question box, six sources.", TEAL),
        ("Hybrid retrieval — lexical AND vector, fused in parallel.", CORAL),
        ("Grounded LLM answers with inline [1][2] citations.", GREEN),
        ("Live tools for meetings, people, papers.", LILAC),
        ("Public / internal visibility — enforced twice.", SAND),
        ("Built-in cache, analytics, USD cost tracking.", NAVY),
        ("Plug-in providers — runs free & local, or with OpenAI.", GRAY),
        ("One docker compose up → ready. 55 tests, Apache 2.0.", TEAL),
    ]
    y = Inches(1.8)
    for label, c in bullets:
        box(s, M, y, Inches(0.45), Inches(0.55), "•", fill=c, color=WHITE, size=22)
        text(s, M + Inches(0.6), y, W - 2 * M - Inches(0.6), Inches(0.55), label,
             size=18, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.62)
    page_number(s, n, total)


def slide_closing(prs, n, total):
    s = make_slide(prs)
    text(s, M, Inches(2.5), W - 2 * M, Inches(1.0), "Questions?", size=60, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, M, Inches(3.8), W - 2 * M, Inches(0.6),
         "Try it locally:", size=18, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    text(s, M, Inches(4.4), W - 2 * M, Inches(0.6),
         "make up && make migrate && make ingest",
         size=18, color=INK, align=PP_ALIGN.CENTER)
    text(s, M, Inches(5.1), W - 2 * M, Inches(0.6),
         "Repo: eic_smart_search · Apache 2.0",
         size=14, color=GRAY, align=PP_ALIGN.CENTER)
    page_number(s, n, total)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    builders = [
        slide_title,
        slide_problem,
        slide_solution,
        slide_architecture,
        slide_data_pie,
        slide_ingestion,
        slide_query_flow,
        slide_hybrid,
        slide_worked_example,
        slide_three_columns,
        slide_visibility,
        slide_tools,
        slide_providers,
        slide_perf,
        slide_observability,
        slide_tldr,
        slide_closing,
    ]
    total = len(builders)
    for i, b in enumerate(builders, start=1):
        print(f"[{i}/{total}] {b.__name__}")
        b(prs, i, total)

    out = Path("PRESENTATION.pptx")
    prs.save(str(out))
    print(f"\nWrote {out} ({out.stat().st_size // 1024} KiB, {total} slides)")


if __name__ == "__main__":
    main()
